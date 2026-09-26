# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl"]
# ///
"""Python version of the PhanTichTuDong macro.

Reads DU_LIEU_GOC, finds the check-out date in the stay description, writes PHAN_TICH_AI and KET_QUA_AI.
Macros inside the .xlsm are kept.

    uv run kiemtra_hoadon.py AI_VBA_KIEM_TRA_HOA_DON_LUU_TRU.xlsm
    uv run kiemtra_hoadon.py --selftest

AI fallback (only for rows the rules can't parse): export OPENAI_API_KEY=sk-...
"""
import datetime as dt
import json
import os
import re
import sys
import urllib.request

TOKEN = re.compile(r"\d+|\D+")


def make_date(d, m, y):
    if len(d) > 2 or len(m) > 2 or len(y) not in (2, 4):
        return None
    year = int(y) + (2000 if len(y) == 2 else 0)
    if not 2000 <= year <= 2100:
        return None
    try:
        return dt.date(year, int(m), int(d))
    except ValueError:  # 31/02, month 13, ...
        return None


def extract_checkout(text):
    """Latest valid date in text: dd/mm/yyyy, dd/mm/yy, dd-mm-yyyy, dd.mm.yyyy, '13 Tháng 5 2025'."""
    t = TOKEN.findall(str(text or ""))
    best = None
    for i in range(len(t) - 4):
        if not (t[i].isdigit() and t[i + 2].isdigit() and t[i + 4].isdigit()):
            continue
        sep, sep2 = t[i + 1], t[i + 3]
        if sep in ("/", "-", ".") and sep2 == sep:
            d = make_date(t[i], t[i + 2], t[i + 4])
        elif re.fullmatch(r"th.ng", sep.strip().lower()) and sep2.replace(",", "").strip() == "":
            d = make_date(t[i], t[i + 2], t[i + 4])
        else:
            continue
        if d and (best is None or d > best):
            best = d
    return best


def to_date(v):
    """Cell value -> date. Text is always read as dd/mm/yyyy."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return extract_checkout(v)


def ask_ai(endpoint, key, model, prompt, text):
    body = json.dumps({"model": model, "instructions": prompt, "input": text}).encode()
    req = urllib.request.Request(endpoint, body, {"Content-Type": "application/json", "Authorization": f"Bearer {key}"})
    with urllib.request.urlopen(req, timeout=30) as r:
        resp = json.load(r)
    return "".join(c.get("text", "") for o in resp.get("output", []) for c in o.get("content", []) or []
                   if c.get("type") == "output_text")


def run(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    wb = load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
    src, ana, res, cfg = (wb[n] for n in ("DU_LIEU_GOC", "PHAN_TICH_AI", "KET_QUA_AI", "CAU_HINH"))
    threshold = int(cfg["B4"].value or 1)
    model, endpoint, prompt = cfg["B5"].value, cfg["B6"].value, cfg["B8"].value
    key = os.environ.get("OPENAI_API_KEY", "")

    for ws, first, cols in ((ana, 2, 7), (res, 11, 9)):
        for row in ws.iter_rows(min_row=first, max_row=max(ws.max_row, first), max_col=cols):
            for c in row:
                c.value = None

    WARN, OK, UNKNOWN = "Cảnh báo", "Bình thường", "Không xác định"
    counts = {WARN: 0, OK: 0, UNKNOWN: 0}
    rows = [r for r in src.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    for n, (so_hd, ngay_hd, mst, ten, info, *_) in enumerate(rows, 1):
        note, method = "", "Quy tắc"
        out = extract_checkout(info)
        if out is None and key:
            method = "AI"
            try:
                answer = ask_ai(endpoint, key, model, prompt, str(info))
                out = extract_checkout(answer)
                if out is None:
                    note = f"AI: {answer}"
            except Exception as e:  # network/API errors: row falls back to manual check
                note = f"Không gọi được AI: {e}"
        inv = to_date(ngay_hd)
        delay = None
        if out is None:
            status, note = UNKNOWN, note or "Cần kiểm tra thủ công"
        elif inv is None:
            status, note = UNKNOWN, "Ngày hóa đơn không hợp lệ"
        else:
            delay = (inv - out).days
            status = WARN if delay >= threshold else OK
        counts[status] += 1

        r = 1 + n
        for col, v in enumerate([n, info, out, method, 1 if out and method != "AI" else None,
                                 "Đã xác định" if out else UNKNOWN, note], 1):
            ana.cell(r, col, v)
        r = 10 + n
        for col, v in enumerate([n, so_hd, inv, mst, ten, info, out, delay, status], 1):
            cell = res.cell(r, col, v)
            if isinstance(v, dt.date):
                cell.number_format = "dd/mm/yyyy"
        c = res.cell(r, 9)
        f = c.font
        c.font = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic,
                      color="FFFF0000" if status == WARN else "FF000000")
    for c in ana["C"][1:]:
        c.number_format = "dd/mm/yyyy"

    res["F6"], res["G6"], res["H6"], res["I6"] = len(rows), counts[OK], counts[WARN], counts[UNKNOWN]
    wb.save(path)
    print(f"{len(rows)} hóa đơn: {counts[WARN]} cảnh báo, {counts[OK]} bình thường, {counts[UNKNOWN]} không xác định")


def selftest():
    D = dt.date
    cases = {
        "Dịch vụ đặt phòng #10258710 - phòng 806 (30/04/2025-02/05/2025)": D(2025, 5, 2),
        "Phòng khách sạn từ ngày 13 Tháng 5 2025 đến ngày 15 Tháng 5 2025 VNTRIP2025D7UKI (Phòng 2001)": D(2025, 5, 15),
        "9 phòng 2 khách x 2 đêm (27/10/2025-29/10/2025) Phòng 1604,1605,1704": D(2025, 10, 29),
        "02/07/2026 14:00 - 03/07/2026 12:00.": D(2026, 7, 3),
        "01/07/26 – 03/07/26.": D(2026, 7, 3),
        "Check in: 01/07/2026; Check out: 03/07/2026.": D(2026, 7, 3),
        "Phòng 1604,1605 không có ngày": None,
        "31/02/2025": None,
        "KHONG_XAC_DINH": None,
    }
    for text, want in cases.items():
        assert extract_checkout(text) == want, (text, extract_checkout(text))
    print("OK")


if __name__ == "__main__":
    if sys.argv[1:] == ["--selftest"]:
        selftest()
    elif len(sys.argv) == 2:
        run(sys.argv[1])
    else:
        sys.exit(__doc__)
